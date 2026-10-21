import os
import time
import yaml
import json
import warnings
import smtplib
import pandas as pd
from pathlib import Path
from dotenv import load_dotenv
from datetime import datetime
from textwrap import dedent
from pydantic import BaseModel
from email.message import EmailMessage
from cert_tracker.tools.helpers import *
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from crewai.flow import Flow, listen, start, and_, router, or_
from cert_tracker.crews.cert_tracker_crew.extraction_crew import ExtractionCrew
from cert_tracker.crews.cert_tracker_crew.models import DataReport


warnings.filterwarnings('ignore')

# class ContentState(BaseModel):
#     topic: str = ""
#     outline: str = ""
#     draft: str = ""
#     final_post: str = ""
# class ContentFlow(Flow[ContentState]):

class ExpiryTrackerFlow(Flow):
    def __init__(self, file_data: dict, test_dirpath: str, email_data:dict, onedrive_folder:str, launchctl_log: str, cert_extraction_csv: str = "", days_ahead=30):
        super().__init__()
        self.days_ahead = days_ahead
        self.file_data = file_data #{filename, {'x_certificate', parsed}}
        self.cert_extraction_csv = cert_extraction_csv
        self.test_dirpath = test_dirpath
        self.email_from = email_data['from']
        self.email_to = email_data['to']
        self.onedrive_folder = onedrive_folder
        self.launchctl_log = launchctl_log
        
    @start()
    def load_existing_data(self):
        if self.cert_extraction_csv:
            #TODO - test branch
            self.state["old_df"] = pd.read_csv(self.cert_extraction_csv, index_col="filename")
            self.state["old_df"]['agent_output'] = self.state['old_df']['agent_output'].map(DataReport.model_validate_json)
        else:
            #TODO - test branch
            self.state["old_df"] = pd.DataFrame()
        self.old_filenames = self.state['old_df'].index.to_list()

    @router(load_existing_data)
    def update_df(self):
        self.state['new_filenames'] = list(self.file_data.keys() - self.old_filenames)
        #get files that are not in old df already
        self.state['new_input_files'] = {k: self.file_data[k] for k in self.state['new_filenames']}

        #if new files, create updated dict 
        if self.state['new_filenames']:
            self.union_filenames = self.file_data.keys() & self.old_filenames
            # self.state['updated_filenames'] = {'existing': union_filenames,
            #                                    'new': self.state['new_filenames']}
            print("cert_tracker.csv needs to be updated.")
            return "needs updating"
        else:
            print("cert_tracker.csv is up to date.")
            return "no updating needed"

    @listen("needs updating")
    def field_extraction(self):        
        results = ExtractionCrew().extraction_crew().kickoff_for_each(inputs=self.state['new_input_files'].values())
        results = [r.pydantic for r in results]
        self.state["extraction_results"] = results
        return results

    @listen(field_extraction)
    def save_extraction(self, results):
        new_df = pd.DataFrame({"agent_output": results}, index=self.state['new_filenames'])
                
        if len(self.state["old_df"]):
            union_df = self.state["old_df"].loc[list(self.union_filenames)]
            self.state['updated_df'] = pd.concat([union_df, new_df], ignore_index=False)
        else:
            self.state['updated_df'] = new_df
        self.state['updated_df'].index.name = "filename" #TODO - redundant?

        self.state['updated_df_pydantic'] = self.state['updated_df'].copy()
        #Save file for updating sharepoint expiry date column
        excel_df = self.state['updated_df'].copy()

        excel_df['Expiry_date'] = self.state['updated_df']['agent_output'].map(lambda x: x.certificate_info.expiry_date)
        excel_df.index.name = "Filename" 

        #convert pydantic object to json string
        self.state['updated_df']['agent_output'] = self.state['updated_df']['agent_output'].map(lambda x: x.model_dump_json())
        
        #save the agent output as a json string, not a reg string, so you can access pydantic object later.
        self.state['updated_df'].to_csv(f"{self.test_dirpath}/cert_extraction.csv", index_label="filename")

        #save excel 
        excel_filepath = Path(f"{self.onedrive_folder}/FlowData/cert_expiry_dates.xlsx")
        excel_df[['Expiry_date']].to_excel(Path("../cert_expiry_dates.xlsx"), index=True, sheet_name="Data")
        print("Saved excel file.")

        # #turn excel sheet into Table format
        wb = load_workbook(Path("../cert_expiry_dates.xlsx"))
        ws = wb["Data"]
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName="Table1", ref=ref)
        ws.add_table(table)
        wb.save(excel_filepath)        
        print("Excel sheet transformed to table.")

        #create FlowTrigger 
        trigger_folder = Path(f"{self.onedrive_folder}/FlowTriggers")
        trigger_file = f"Trigger_{datetime.now():%m%d%Y}.txt"

        with open(trigger_folder / trigger_file, "w") as f:
            f.write("GO!")

    @listen(or_ ("no updating needed", save_extraction))
    def get_expiry_dates(self):
        if self.state['new_filenames']:
            self.state['expiry_dates'] = [x.certificate_info.expiry_date for x in self.state['updated_df_pydantic']['agent_output']]
        else:
            type(self.state['old_df']['agent_output'])
            self.state['expiry_dates'] = [x.certificate_info.expiry_date for x in self.state['old_df']['agent_output']] 
    
    #get docs that are close to expiry only
    @router(get_expiry_dates)
    def filter_by_expiry(self):
        person_names = []
        expired_dates = []
        course_names = []
        expiry_filenames = []
        
        if "updated_df" not in self.state:
            df = self.state["old_df"]
        else:
            df = self.state['updated_df_pydantic']

        for i,r in enumerate(df['agent_output']):
            expiry_date = self.state['expiry_dates'][i]
            date_diff = get_expired(expiry_date, self.days_ahead) #TODO - edge case fail - llm fails to follow date format
    
            if date_diff:
                course_names.append(r.certificate_info.course_name) 
                person_names.append(r.personal_info.name)
                expired_dates.append(expiry_date)
                expiry_filenames.append(df.index[i])
    
        expired_cases = pd.DataFrame({'cert_owner': person_names,
                                     'course_name': course_names,
                                     'expiry_date': expired_dates}, 
                                     index= expiry_filenames)
        expired_cases.index.name = "filename"

        self.state['curr_expired_df'] = expired_cases
        
        if expired_cases.empty:
            self.state['curr_expired_df'].to_csv(f"{self.test_dirpath}/cert_expiring.csv", index_label="filename")
            print("No expiring cases")
            return 'No expiring cases'
        elif len(expired_cases) == 1:
            self.state['curr_expired_df'].to_csv(f"{self.test_dirpath}/cert_expiring.csv", index_label="filename")
            return 'One expiring case'
        else:
            return 'Multiple expiring'

    @listen('Multiple expiring')
    def ordered_expiry(self):
        self.state['curr_expired_df'] = self.state['curr_expired_df'].sort_values('expiry_date')
        self.state['curr_expired_df'].to_csv(f"{self.test_dirpath}/cert_expiring.csv", index_label="filename")

    @listen(or_ ("One expiring case", ordered_expiry))
    def email_content(self):
        cert_info = format_certificates(self.state['curr_expired_df'])
        self.subject = f"""Certificates expiring within the next {self.days_ahead} days"""

        template = dedent("""\
        Dear {recipient},

        I am writing to let you know the following certificates are due to expire within the next {days} days:

        Certificates:

        {certificates}

        Kind regards,
        {sender}
        """)

        self.content = template.format(
            recipient = self.email_to['name'],
            days = self.days_ahead,
            certificates = cert_info,
            sender = self.email_from['name']
        )

    @router(email_content)
    def email_due(self):
        STATE = Path.home() / self.launchctl_log 
        if is_due(STATE):
            STATE.write_text(date.today().isoformat())
            return "Email is due"
        else:
            print("Email notification is not due yet.")
            return "Email not due yet"

    #send email
    @listen('Email is due')
    def send_email(self):
        msg = EmailMessage()
        msg["Subject"] = self.subject
        msg["From"] = self.email_from['email']
        msg["To"] = self.email_to['email']
        msg.set_content(self.content)
        
        with smtplib.SMTP("smtp.gmail.com", 587) as server:
            server.starttls()
            server.ehlo()
            server.login(self.email_from['email'], os.getenv("GMAIL_PASS"))
            server.send_message(msg)


def kickoff(file_data, test_dirpath, email_addresses, onedrive_folder, launchctl_log, cert_extraction_csv, days_ahead):
    flow = ExpiryTrackerFlow(file_data=file_data, test_dirpath=test_dirpath, days_ahead=days_ahead, onedrive_folder=onedrive_folder, launchctl_log=launchctl_log, cert_extraction_csv=cert_extraction_csv, email_data=email_data, ) 
    flow.kickoff()

if __name__ == "__main__":
    with open("cert_tracker/crews/cert_tracker_crew/config/constants.yaml", "r") as f:
        constants_config = yaml.safe_load(f)
    gcp_files = get_all_filepaths(constants_config['gcp_foldername'])
    div_files = get_all_filepaths(constants_config['div_foldername'])
    tdg_files = get_all_filepaths(constants_config['tdg_foldername'])
    email_data = {'from': {'name': constants_config['email_from']['name'],
                            'email': constants_config['email_from']['email']},
                    'to': {'name': constants_config['email_to']['name'],
                            'email': constants_config['email_to']['email']}}
    file_data = parsed_files(gcp_filepath = gcp_files, tdg_filepath = tdg_files, div_filepath = div_files) 

    #load variables from .env
    load_dotenv()

    start = time.perf_counter()
    kickoff(file_data, constants_config['test_dirpath'], email_data, constants_config['one_drive_folder'], constants_config['launchctl_log'], constants_config['cert_extraction_csv'], constants_config['days_ahead'])
    elapsed = time.perf_counter() - start
    print(f"flow took {elapsed/ 60:.2f} minutes")